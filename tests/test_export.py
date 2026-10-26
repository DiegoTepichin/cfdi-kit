from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from cfdi_kit.export import (
    INVOICE_COLUMNS,
    Direction,
    concepts_frame,
    direction,
    invoices_frame,
    payments_frame,
    to_csv,
    to_excel,
)
from cfdi_kit.models import Cfdi
from tests.conftest import MY_RFC, load


def test_direction_is_relative_to_the_rfc() -> None:
    assert direction(load("income_pue_40.xml"), MY_RFC) is Direction.ISSUED
    assert direction(load("expense_pue_40.xml"), MY_RFC.lower()) is Direction.RECEIVED
    assert direction(load("payroll_40.xml"), MY_RFC) is Direction.OTHER


def test_invoices_frame_has_one_row_per_cfdi(all_valid: list[Cfdi]) -> None:
    frame = invoices_frame(all_valid, MY_RFC)

    assert list(frame.columns) == list(INVOICE_COLUMNS)
    assert len(frame) == len(all_valid)
    assert pd.api.types.is_datetime64_any_dtype(frame["issued_at"])


def test_invoices_frame_values() -> None:
    row = invoices_frame([load("income_pue_40.xml")], MY_RFC).iloc[0]

    assert row["direction"] == "issued"
    assert row["type_label"] == "Ingreso"
    assert row["payment_form_label"] == "Transferencia electrónica de fondos"
    assert row["iva_transferred"] == pytest.approx(1600.0)
    assert row["isr_withheld"] == pytest.approx(1000.0)
    assert row["total_mxn"] == pytest.approx(9533.33)


def test_total_mxn_uses_the_exchange_rate() -> None:
    row = invoices_frame([load("expense_ppd_usd_33.xml")]).iloc[0]

    assert row["direction"] is None
    assert row["total_mxn"] == pytest.approx(290 * 18.5)


def test_related_uuids_are_listed() -> None:
    row = invoices_frame([load("credit_note_40.xml")]).iloc[0]

    assert row["related_uuids"] == "6F1C2A7E-1B2C-4D3E-9F40-0A1B2C3D4E01"


def test_empty_input_keeps_the_columns() -> None:
    assert list(invoices_frame([]).columns) == list(INVOICE_COLUMNS)
    assert payments_frame([]).empty


def test_concepts_frame() -> None:
    frame = concepts_frame([load("income_pue_40.xml"), load("expense_pue_40.xml")])

    assert len(frame) == 2
    assert frame["iva_withheld"].tolist() == pytest.approx([1066.67, 0.0])
    assert frame["discount"].tolist() == pytest.approx([0.0, 100.0])


def test_payments_frame() -> None:
    frame = payments_frame([load("payment_20.xml"), load("payment_10.xml")])

    assert frame["document_uuid"].tolist() == [
        "6F1C2A7E-1B2C-4D3E-9F40-0A1B2C3D4E02",
        "6F1C2A7E-1B2C-4D3E-9F40-0A1B2C3D4E07",
    ]
    assert frame["iva_paid"].iloc[0] == pytest.approx(3200.0)
    assert pd.isna(frame["iva_paid"].iloc[1])


def test_to_csv_round_trips(tmp_path: Path, all_valid: list[Cfdi]) -> None:
    path = tmp_path / "out.csv"

    to_csv(all_valid, path, MY_RFC)

    assert len(pd.read_csv(path)) == len(all_valid)


def test_to_excel_writes_three_formatted_sheets(tmp_path: Path, all_valid: list[Cfdi]) -> None:
    path = tmp_path / "out.xlsx"

    to_excel(all_valid, path, MY_RFC)

    workbook = load_workbook(path)
    assert workbook.sheetnames == ["invoices", "line_items", "payments"]
    sheet = workbook["invoices"]
    assert sheet.max_row == len(all_valid) + 1
    assert sheet.freeze_panes == "A2"
    assert sheet["A1"].value == "uuid"
